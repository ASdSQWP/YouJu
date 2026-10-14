"""把上传的文件变成带章节和页码的文本块。"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from xml.etree import ElementTree as ET

SUPPORTED = (".md", ".markdown", ".txt", ".pdf", ".docx", ".html", ".htm", ".csv", ".xlsx")


class ParseError(Exception):
    pass


@dataclass
class Block:
    text: str
    section: str = ""
    page: int | None = None


def decode_bytes(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _join(path: list[str]) -> str:
    return " / ".join(p for p in path if p)[:200]


_MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")


def parse_markdown(text: str) -> list[Block]:
    lines = text.splitlines()
    h1_count = sum(1 for line in lines if re.match(r"^#\s+\S", line))
    skip_title = h1_count == 1
    blocks: list[Block] = []
    path: list[str] = []
    buf: list[str] = []
    in_code = False

    def flush() -> None:
        body = "\n".join(buf).strip()
        if body:
            visible = path[1:] if skip_title and path else path
            blocks.append(Block(body, _join(visible)))
        buf.clear()

    for line in lines:
        if line.strip().startswith("```"):
            in_code = not in_code
            buf.append(line)
            continue
        if not in_code:
            match = _MD_HEADING.match(line)
            if match:
                flush()
                level = len(match.group(1))
                path = (path + [""] * level)[: level - 1] + [match.group(2).strip()]
                continue
            if not line.strip():
                flush()
                continue
        buf.append(line)
    flush()
    return blocks


_TXT_HEADING = re.compile(r"^(第[一二三四五六七八九十百零\d]+[章节条篇部分]|[一二三四五六七八九十]+[、．.]|【[^】]{1,20}】)")


def _looks_heading(line: str) -> bool:
    return len(line) <= 30 and not line.endswith(("。", "；", "，", ";", ",", "：", ":")) and bool(_TXT_HEADING.match(line))


def parse_plain(text: str, section: str = "", page: int | None = None) -> tuple[list[Block], str]:
    blocks: list[Block] = []
    buf: list[str] = []

    def flush() -> None:
        body = "\n".join(buf).strip()
        if body:
            blocks.append(Block(body, section, page))
        buf.clear()

    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            flush()
            continue
        if _looks_heading(line):
            flush()
            section = line
            continue
        buf.append(line)
    flush()
    return blocks, section


def _reflow(text: str) -> str:
    """PDF 抽出来的文字按版面硬换行，把句子中间的换行接回去。"""
    out: list[str] = []
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            out.append("")
            continue
        if out and out[-1] and not out[-1].endswith(("。", "！", "？", "；", "：", ".", "!", "?", ";", ":")) and not re.match(r"^([（(]?\d+[)）.、]|[-•●]|第.+[章节条])", line):
            prev = out[-1]
            joiner = " " if prev[-1:].isascii() and line[:1].isascii() else ""
            out[-1] = prev + joiner + line
        else:
            out.append(line)
    return "\n".join(out)


def parse_pdf(path: Path) -> list[Block]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ParseError("缺少 pypdf，无法解析 PDF") from exc
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            reader.decrypt("")
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"PDF 无法打开：{exc}") from exc
    blocks: list[Block] = []
    section = ""
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001
            text = ""
        found, section = parse_plain(_reflow(text), section, number)
        blocks.extend(found)
    if not blocks:
        raise ParseError("PDF 中没有可提取的文字，可能是扫描件，需要先做 OCR")
    return blocks


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx_styles(zf: zipfile.ZipFile) -> dict[str, int]:
    try:
        root = ET.fromstring(zf.read("word/styles.xml"))
    except (KeyError, ET.ParseError):
        return {}
    levels: dict[str, int] = {}
    for style in root.findall(f"{_W}style"):
        style_id = style.get(f"{_W}styleId") or ""
        name_el = style.find(f"{_W}name")
        name = (name_el.get(f"{_W}val") if name_el is not None else "") or ""
        match = re.match(r"(?:heading|标题)\s*(\d)", name.strip().lower())
        if match:
            levels[style_id] = int(match.group(1))
        elif name.strip().lower() == "title":
            levels[style_id] = 1
    return levels


def _para_text(el) -> str:
    parts: list[str] = []
    for node in el.iter():
        if node.tag == f"{_W}t" and node.text:
            parts.append(node.text)
        elif node.tag == f"{_W}tab":
            parts.append("\t")
        elif node.tag in (f"{_W}br", f"{_W}cr"):
            parts.append("\n")
    return "".join(parts)


def _heading_level(el, styles: dict[str, int]) -> int:
    ppr = el.find(f"{_W}pPr")
    if ppr is None:
        return 0
    style = ppr.find(f"{_W}pStyle")
    if style is not None:
        level = styles.get(style.get(f"{_W}val") or "", 0)
        if level:
            return level
    outline = ppr.find(f"{_W}outlineLvl")
    if outline is not None:
        try:
            value = int(outline.get(f"{_W}val") or "9")
        except ValueError:
            return 0
        if value < 6:
            return value + 1
    return 0


def parse_docx(path: Path) -> list[Block]:
    try:
        with zipfile.ZipFile(path) as zf:
            xml = zf.read("word/document.xml")
            styles = _docx_styles(zf)
    except (zipfile.BadZipFile, KeyError) as exc:
        raise ParseError("Word 文件损坏，或不是 .docx 格式") from exc
    body = ET.fromstring(xml).find(f"{_W}body")
    if body is None:
        return []
    blocks: list[Block] = []
    path_titles: list[str] = []
    for el in body:
        if el.tag == f"{_W}p":
            text = _para_text(el).strip()
            if not text:
                continue
            level = _heading_level(el, styles)
            if level:
                path_titles = (path_titles + [""] * level)[: level - 1] + [text]
                continue
            blocks.append(Block(text, _join(path_titles)))
        elif el.tag == f"{_W}tbl":
            rows = []
            for tr in el.iter(f"{_W}tr"):
                cells = [_para_text(tc).strip() for tc in tr.iter(f"{_W}tc")]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                blocks.append(Block("\n".join(rows), _join(path_titles)))
    return blocks


class _HTMLToMarkdown(HTMLParser):
    SKIP = {"script", "style", "noscript", "head", "svg"}
    BLOCK = {"p", "div", "section", "article", "tr", "br", "ul", "ol", "table", "blockquote", "header", "footer"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.out: list[str] = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip += 1
        elif re.fullmatch(r"h[1-6]", tag):
            self.out.append("\n\n" + "#" * int(tag[1]) + " ")
        elif tag == "li":
            self.out.append("\n- ")
        elif tag in ("td", "th"):
            self.out.append(" | ")
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self.skip = max(0, self.skip - 1)
        elif re.fullmatch(r"h[1-6]", tag) or tag in ("p", "table"):
            self.out.append("\n\n")

    def handle_data(self, data):
        if not self.skip:
            self.out.append(data.replace("\n", " "))


def parse_html(text: str) -> list[Block]:
    parser = _HTMLToMarkdown()
    parser.feed(text)
    markdown = re.sub(r"[ \t]+", " ", "".join(parser.out))
    return parse_markdown(markdown)


def _rows_to_blocks(rows: list[list[str]], section: str) -> list[Block]:
    rows = [[str(cell).strip() for cell in row] for row in rows if any(str(cell).strip() for cell in row)]
    if not rows:
        return []
    header = rows[0]
    if len(rows) == 1:
        return [Block(" | ".join(header), section)]
    blocks = []
    for row in rows[1:]:
        pairs = []
        for i, cell in enumerate(row):
            if not cell:
                continue
            name = header[i] if i < len(header) and header[i] else f"列{i + 1}"
            pairs.append(f"{name}：{cell}")
        if pairs:
            blocks.append(Block("；".join(pairs), section))
    return blocks


def parse_xlsx(path: Path) -> list[Block]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ParseError("缺少 openpyxl，无法解析 Excel") from exc
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"Excel 无法打开：{exc}") from exc
    blocks: list[Block] = []
    for sheet in workbook.worksheets:
        rows = [["" if value is None else value for value in row] for row in sheet.iter_rows(values_only=True)]
        blocks.extend(_rows_to_blocks(rows, sheet.title))
    workbook.close()
    return blocks


def parse_file(path: Path, filename: str) -> list[Block]:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise ParseError(f"不支持的文件类型：{ext or '无扩展名'}")
    if ext == ".pdf":
        return parse_pdf(path)
    if ext == ".docx":
        return parse_docx(path)
    if ext == ".xlsx":
        return parse_xlsx(path)
    text = decode_bytes(path.read_bytes())
    if ext in (".md", ".markdown"):
        return parse_markdown(text)
    if ext in (".html", ".htm"):
        return parse_html(text)
    if ext == ".csv":
        return _rows_to_blocks(list(csv.reader(io.StringIO(text))), Path(filename).stem)
    return parse_plain(text)[0]
